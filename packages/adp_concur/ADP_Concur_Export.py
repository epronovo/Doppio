"""
ADP_Concur_Export - write the flat file SAP Concur picks up.

One file carries all three record types. Each line is one record: the type in
the first field, then the template's columns in order, every position present
whether or not it has a value - which is what "All fields must be represented"
on the 350 and 360 tabs means.

Two orderings are offered because sites differ on this. 'by_type' writes every
305, then every 350, then every 360, which is the safer default: Concur creates
the profile from the 305, and the 350 and 360 records attach to a profile that
has to exist already. 'by_employee' writes a person's three records together.

Nothing here decides who is in the file - ADP_Concur_Map.selected_employees()
does, from the include flags, the configured scope and the exception list.
"""
from __future__ import annotations

import argparse
import csv
import io
import sqlite3
from datetime import datetime
from pathlib import Path

from ADP_Concur_Db import DEFAULT_DB_PATH, ROSTERS, connect, resolve_db_path
from ADP_Concur_Map import (
    build_import_settings,
    build_record,
    layout_width,
    load_config,
    selected_employees,
)

BASE_DIR = Path(__file__).parent.resolve()
DEFAULT_OUTPUT_DIR = BASE_DIR / "output" / "adp_concur"

QUOTING = {"minimal": csv.QUOTE_MINIMAL, "all": csv.QUOTE_ALL, "none": csv.QUOTE_NONE}


def outbound_dir(cfg: dict) -> Path:
    """Where the file lands - the configured folder, or output/adp_concur/."""
    configured = (cfg.get("extract") or {}).get("outbound_dir") or ""
    return Path(configured).expanduser() if configured else DEFAULT_OUTPUT_DIR


def extract_file_name(cfg: dict, when: datetime | None = None,
                      selection: bool = False, roster: str = "",
                      kind: str = "") -> str:
    """
    {stamp} in the configured name becomes yyyymmdd_HHMMSS.

    A selection gets its own pattern, because the outbound folder is a Concur
    pickup and a partial file that looks exactly like a full one is the kind of
    thing that gets loaded by accident at four in the afternoon. `kind="320"`
    does the same for the Update ID file, which must never be mistaken for the
    305/350/360 file sitting in the same folder.
    """
    when = when or datetime.now()
    extract = cfg.get("extract") or {}
    if kind == "320":
        pattern = extract.get("file_name_320") or "FMG_Concur_UpdateID_{stamp}.txt"
    else:
        pattern = (extract.get("selection_file_name")
                   or "FMG_Concur_Employee_Selection_{stamp}.txt") if selection else (
            extract.get("file_name") or "FMG_Concur_Employee_{stamp}.txt")
    name = pattern.format(stamp=when.strftime("%Y%m%d_%H%M%S"),
                          date=when.strftime("%Y%m%d"),
                          time=when.strftime("%H%M%S"))
    if roster:
        # The two rosters are two separate loads into Concur and land in the
        # same pickup folder seconds apart, so the file has to say which is
        # which - a stamp alone would not.
        stem, dot, ext = name.rpartition(".")
        name = f"{stem or name}_{roster.upper()}{dot}{ext}" if dot else \
            f"{name}_{roster.upper()}"
    return name


def record_types_for(cfg: dict, roster: str = "") -> list[str]:
    """
    Which record types one roster's file carries, in writing order.

    Config first, then the roster's built-in list. In config because it is a
    real question rather than a constant: the July file that loaded carried 305
    and 360 only, and the workbook has since grown a 350 tab.

    records_enabled is layered on top of that and applies whether or not a
    roster is given, so it is the one place that can turn a type off for
    every file - including the plain combined export, which has no roster to
    look up in records_by_roster at all.
    """
    wanted = (["305", "350", "360"] if roster not in ROSTERS else
              (cfg.get("records_by_roster") or {}).get(roster)
              or ROSTERS[roster]["records"])
    enabled = cfg.get("records_enabled") or {}
    return [t for t in ("305", "350", "360")
            if t in wanted and enabled.get(t, True)]


def collect(conn: sqlite3.Connection, cfg: dict,
            keys: list[int] | None = None,
            roster: str = "") -> tuple[list[list[str]], dict]:
    """
    Every record the extract will carry, in the configured order.

    Returns the lines and a per-type count. Built here rather than in
    ADP_Concur_Map.build_records() because the ordering is an export concern.

    The file opens with the 100 Import Settings record, which SAP requires once
    per file and which is not per-employee - so it is prepended here rather
    than living in a field map with the rest.

    `roster` narrows it to one population and, with it, to that population's
    record types: the non-US roster has no 350 tab in the workbook and gets no
    350 records.
    """
    cfg_extract = cfg.get("extract") or {}
    types = record_types_for(cfg, roster)
    widths = {rt: layout_width(conn, rt) for rt in types}
    people = {rt: selected_employees(conn, rt, cfg, keys, roster) for rt in types}
    # Reported for all three types whatever the roster carries, so a caller
    # never has to know which types a roster has to read the numbers.
    counts = {rt: len(people.get(rt, ())) for rt in ("305", "350", "360")}

    lines: list[list[str]] = [build_import_settings(cfg)]
    if cfg_extract.get("order") == "by_employee":
        # Keyed on file number so the three records for one person stay
        # together; the 305 order is the one that drives the file.
        wanted = {rt: {e["file_number"]: e for e in people[rt]} for rt in people}
        seen = []
        for rt in types:
            for e in people[rt]:
                if e["file_number"] not in seen:
                    seen.append(e["file_number"])
        for fn in seen:
            for rt in types:
                emp = wanted[rt].get(fn)
                if emp is not None:
                    lines.append(build_record(emp, rt, widths[rt], cfg))
    else:
        for rt in types:
            for emp in people[rt]:
                lines.append(build_record(emp, rt, widths[rt], cfg))

    return lines, counts


def render(lines: list[list[str]], cfg: dict) -> str:
    """The file as text, so it can be previewed without being written."""
    extract = cfg.get("extract") or {}
    buf = io.StringIO()
    writer = csv.writer(
        buf,
        delimiter=extract.get("delimiter", ","),
        quoting=QUOTING.get(extract.get("quote", "minimal"), csv.QUOTE_MINIMAL),
        quotechar='"',
        escapechar="\\" if extract.get("quote") == "none" else None,
        lineterminator=extract.get("line_ending", "\r\n"),
    )
    writer.writerows(lines)
    return buf.getvalue()


def ADP_Concur_export(conn: sqlite3.Connection, cfg: dict | None = None,
                      out_dir: str | Path | None = None,
                      file_name: str | None = None,
                      keys: list[int] | None = None,
                      selection_label: str = "",
                      roster: str = "",
                      dry_run: bool = False) -> dict:
    """
    Write the extract.

    `keys` narrows it to a chosen set of people - one manager's organisation,
    a filtered list, whatever was ticked. The file is named and recorded as a
    selection so a pilot load is never mistaken for the full company, and
    `selection_label` is stored alongside it saying what was chosen.

    dry_run builds everything and reports the counts without touching the
    disk, which is what the front end previews with.
    """
    cfg = cfg or load_config()
    lines, counts = collect(conn, cfg, keys, roster)
    text = render(lines, cfg)

    directory = Path(out_dir).expanduser() if out_dir else outbound_dir(cfg)
    name = file_name or extract_file_name(cfg, selection=keys is not None,
                                          roster=roster)
    target = directory / name
    scope = ",".join(f"{k}={v}" for k, v in (cfg.get("scope") or {}).items())
    if roster:
        scope = f"roster {roster}; " + scope
    if keys is not None:
        scope = (f"selection of {len(keys)}"
                 + (f" ({selection_label})" if selection_label else "")
                 + "; " + scope)

    result = {"file_name": name, "path": str(target), "records": len(lines),
              "n_305": counts["305"], "n_350": counts["350"],
              "n_360": counts["360"], "bytes": len(text.encode(
                  (cfg.get("extract") or {}).get("encoding", "utf-8"))),
              "dry_run": dry_run, "scope": scope, "roster": roster,
              "roster_label": ROSTERS.get(roster, {}).get("label", "Everyone"),
              "record_types": record_types_for(cfg, roster),
              "selection": keys is not None,
              "selected": len(keys) if keys is not None else None,
              "selection_label": selection_label,
              "preview": text.splitlines()[:5]}
    if dry_run:
        return result

    directory.mkdir(parents=True, exist_ok=True)
    encoding = (cfg.get("extract") or {}).get("encoding", "utf-8")
    with open(target, "w", encoding=encoding, newline="") as fh:
        fh.write(text)

    conn.execute(
        "INSERT INTO ADP_Concur_Extracts "
        "(file_name, file_path, n_305, n_350, n_360, scope, delimiter) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (name, str(target), counts["305"], counts["350"], counts["360"], scope,
         (cfg.get("extract") or {}).get("delimiter", ",")))
    conn.commit()
    return result


def ADP_Concur_export_all(conn: sqlite3.Connection, cfg: dict | None = None,
                          out_dir: str | Path | None = None,
                          dry_run: bool = False) -> list[dict]:
    """
    One file per roster - the two loads Concur actually wants.

    They cannot be one file: the two populations have different record types
    and, in the workbook, different Login ID rules, and Concur takes one 100
    record per file. So each roster gets its own file with its own 100 record
    on the front, and a roster with nobody in it is skipped rather than written
    as an empty file.
    """
    cfg = cfg or load_config()
    out = []
    for roster in ROSTERS:
        n = conn.execute(
            "SELECT COUNT(*) FROM ADP_Concur_Employees "
            "WHERE roster = ? AND row_state <> 'deleted'", (roster,)).fetchone()[0]
        if not n:
            continue
        out.append(ADP_Concur_export(conn, cfg, out_dir=out_dir, roster=roster,
                                     dry_run=dry_run))
    return out


def ADP_Concur_export_320(conn: sqlite3.Connection, cfg: dict | None = None,
                          out_dir: str | Path | None = None,
                          file_name: str | None = None,
                          keys: list[int] | None = None,
                          selection_label: str = "",
                          dry_run: bool = False) -> dict:
    """
    Write the standalone 320 file - Update ID Information.

    Deliberately not a record type ADP_Concur_export() ever writes: SAP
    requires the 320 be uploaded separately from the 305/310, a day ahead, and
    never merged into the same file - see the note on Table 6 of the Employee
    Import Specification. It also has no roster - Login ID applies either side
    of the US / non-US line.

    include_320 is on by default like the 305/350/360, because SAP's own
    guidance is to carry Login ID changes through the 320 rather than the 305
    - see the field map's comment in ADP_Concur_Map.FIELD_MAP. It stays
    per-employee so any one person can still be held out of this file without
    touching anything else, the same as the other three record types.
    """
    cfg = cfg or load_config()
    width = layout_width(conn, "320")
    people = selected_employees(conn, "320", cfg, keys)
    lines = [build_import_settings(cfg)] + [
        build_record(e, "320", width, cfg) for e in people]
    text = render(lines, cfg)

    directory = Path(out_dir).expanduser() if out_dir else outbound_dir(cfg)
    name = file_name or extract_file_name(cfg, selection=keys is not None, kind="320")
    target = directory / name
    scope = f"320={(cfg.get('scope') or {}).get('320', 'all')}"
    if keys is not None:
        scope = (f"selection of {len(keys)}"
                 + (f" ({selection_label})" if selection_label else "")
                 + "; " + scope)

    result = {"file_name": name, "path": str(target), "records": len(lines),
              "n_320": len(people), "bytes": len(text.encode(
                  (cfg.get("extract") or {}).get("encoding", "utf-8"))),
              "dry_run": dry_run, "scope": scope,
              "selection": keys is not None,
              "selected": len(keys) if keys is not None else None,
              "selection_label": selection_label,
              "preview": text.splitlines()[:5]}
    if dry_run:
        return result

    directory.mkdir(parents=True, exist_ok=True)
    encoding = (cfg.get("extract") or {}).get("encoding", "utf-8")
    with open(target, "w", encoding=encoding, newline="") as fh:
        fh.write(text)

    conn.execute(
        "INSERT INTO ADP_Concur_Extracts "
        "(file_name, file_path, n_320, scope, delimiter) VALUES (?, ?, ?, ?, ?)",
        (name, str(target), len(people), scope,
         (cfg.get("extract") or {}).get("delimiter", ",")))
    conn.commit()
    return result


def held_back(conn: sqlite3.Connection,
              keys: list[int] | None = None) -> list[dict]:
    """
    Who the extract is leaving out, and why.

    Anyone with a blocking error, plus anyone excluded by an include flag.
    Shown beside the export button so the file is never quietly short. With a
    selection, only people inside that selection are reported - the other
    hundred are not being "held back", they were simply not chosen.
    """
    scope, args = "", []
    if keys is not None:
        if not keys:
            return []
        scope = f" AND e.employee_key IN ({','.join('?' * len(keys))})"
        args = list(keys)
    return [dict(r) for r in conn.execute(
        f"""
        SELECT e.file_number,
               e.legal_last_name || ', ' || e.legal_first_name AS name,
               e.position_status, e.concur_status,
               e.include_305, e.include_350, e.include_360,
               (SELECT GROUP_CONCAT(x.message, ' | ') FROM ADP_Concur_Exceptions x
                 WHERE x.employee_key = e.employee_key AND x.severity = 'error')
                   AS reasons
          FROM ADP_Concur_Employees e
         WHERE e.row_state <> 'deleted'
           AND (EXISTS (SELECT 1 FROM ADP_Concur_Exceptions x
                         WHERE x.employee_key = e.employee_key AND x.severity = 'error')
                OR e.include_305 = 0 OR e.include_350 = 0 OR e.include_360 = 0)
               {scope}
         ORDER BY e.legal_last_name, e.legal_first_name
        """, args)]


def ADP_Concur_export_exceptions(rows: list[dict]) -> io.BytesIO:
    """
    The Exceptions list as an .xlsx workbook, for review outside the browser.

    Same rows the Exceptions tab shows - whatever severity/field/picked filter
    was in effect when the download was asked for. One sheet, header frozen,
    errors filled red and warnings amber so the sheet reads the same way the
    tab's pills do.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    columns = [("severity", "Severity"), ("file_number", "File #"),
              ("employee_name", "Name"), ("field", "Check"),
              ("message", "What is wrong"), ("position_status", "Status"),
              ("concur_status", "Concur active"),
              ("business_unit_desc", "Business unit"), ("source", "Source")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    error_fill = PatternFill("solid", fgColor="FBE2E2")
    warning_fill = PatternFill("solid", fgColor="FDF3D9")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(key, "") for key, _ in columns])
        fill = error_fill if row.get("severity") == "error" else warning_fill
        for cell in ws[ws.max_row]:
            cell.fill = fill
        ws.cell(ws.max_row, 5).alignment = wrap

    widths = [10, 10, 24, 20, 60, 12, 14, 24, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Write the Concur employee extract.")
    ap.add_argument("--db", default=None, help=f"SQLite path (default {DEFAULT_DB_PATH})")
    ap.add_argument("--out-dir", default=None,
                    help=f"Where to write (default {DEFAULT_OUTPUT_DIR})")
    ap.add_argument("--name", default=None, help="Override the file name")
    ap.add_argument("--dry-run", action="store_true",
                    help="Report what would be written without writing it")
    ap.add_argument("--held-back", action="store_true",
                    help="List the people the extract leaves out and why")
    ap.add_argument("--under", metavar="FILE_NUMBER", action="append", default=[],
                    help="Scope the file to this person and everyone under "
                         "them. Repeatable - the perfect pilot load")
    ap.add_argument("--only", metavar="FILE_NUMBER", action="append", default=[],
                    help="Scope the file to these people exactly. Repeatable")
    ap.add_argument("--roster", choices=list(ROSTERS),
                    help="Write one roster's file only")
    ap.add_argument("--all-rosters", action="store_true",
                    help="Write one file per roster - the two loads Concur wants")
    args = ap.parse_args(argv)

    conn = connect(args.db)
    print(f"Database: {resolve_db_path(args.db)}")

    keys = None
    label = ""
    if args.under or args.only:
        from ADP_Concur_Hierarchy import ADP_Concur_subtree_keys
        keys = ADP_Concur_subtree_keys(conn, args.under) if args.under else []
        if args.only:
            marks = ",".join("?" * len(args.only))
            keys += [r[0] for r in conn.execute(
                f"SELECT employee_key FROM ADP_Concur_Employees "
                f"WHERE file_number IN ({marks}) AND row_state <> 'deleted'",
                args.only) if r[0] not in keys]
        label = "; ".join(filter(None, [
            "under " + ", ".join(args.under) if args.under else "",
            "only " + ", ".join(args.only) if args.only else ""]))
        print(f"Selection: {label} — {len(keys)} employee(s)")

    if args.held_back:
        rows = held_back(conn, keys)
        print(f"{len(rows)} employee(s) held back")
        for r in rows:
            flags = "".join(t for t, k in (("305", "include_305"), ("350", "include_350"),
                                           ("360", "include_360")) if not r[k])
            print(f"  {r['file_number']:<8} {(r['name'] or ''):<30} "
                  f"{r['reasons'] or ('excluded from ' + flags)}")
        conn.close()
        return 0

    if args.all_rosters:
        for res in ADP_Concur_export_all(conn, out_dir=args.out_dir,
                                         dry_run=args.dry_run):
            verb = "Would write" if args.dry_run else "Wrote"
            print(f"{verb} {res['roster_label']}: {res['records']} record(s) "
                  f"(1 x 100, {res['n_305']} x 305, {res['n_350']} x 350, "
                  f"{res['n_360']} x 360)")
            print(f"  {res['path']}  ({res['bytes']:,} bytes)")
        conn.close()
        return 0

    res = ADP_Concur_export(conn, out_dir=args.out_dir, file_name=args.name,
                            keys=keys, selection_label=label,
                            roster=args.roster or "", dry_run=args.dry_run)
    verb = "Would write" if args.dry_run else "Wrote"
    print(f"{verb} {res['records']} record(s): 1 x 100, {res['n_305']} x 305, "
          f"{res['n_350']} x 350, {res['n_360']} x 360")
    print(f"  {res['path']}  ({res['bytes']:,} bytes)")
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
