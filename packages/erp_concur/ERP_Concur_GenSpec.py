"""One-off generator: Concur_Record_Type_Specifications.xlsx -> ERP_Concur_Spec.py."""
import datetime as dt
from pathlib import Path
import openpyxl

WB = Path(__file__).resolve().parent.parent.parent / "FMG" / "Concur_Record_Type_Specifications.xlsx"
OUT = Path(__file__).resolve().parent / "ERP_Concur_Spec.py"

# tab -> (spec key, record type, the app's short label)
TABS = {
    "AV v3 - 100 Import Settings": ("vendor_100", "100", "Import Settings"),
    "AV v3 - 200 Vendor":          ("vendor_200", "200", "Vendor"),
    "PO - 200 Request Header":     ("po_200",     "200", "Request Header"),
    "PO - 210 Bill-to Address":    ("po_210",     "210", "Bill-to Address"),
    "PO - 220 Ship-to Address":    ("po_220",     "220", "Ship-to Address"),
    "PO - 300 Line Item":          ("po_300",     "300", "Line Item"),
    "PO - 400 Line Allocation":    ("po_400",     "400", "Line Allocation"),
    "PO Receipt v2 - 200 Header":  ("receipt_200","200", "Receipt"),
}

wb = openpyxl.load_workbook(WB, data_only=True)
spec = {}
for tab, (key, rt, label) in TABS.items():
    ws = wb[tab]
    header = [ws.cell(7, c).value for c in range(1, 14)]
    col_req = header.index("Required?") + 1
    col_proc = header.index("Mapped From (proc)") + 1
    fields = []
    r = 8
    while ws.cell(r, 1).value not in (None, ""):
        name = str(ws.cell(r, 1).value).strip()
        req = ws.cell(r, col_req).value
        proc = ws.cell(r, col_proc).value
        # Section banners ("Vendor Header", "Vendor Address Detail") are the
        # only rows with nothing in any other column - a real field always
        # carries at least a definition. Blank Required? does NOT mark a
        # banner: four vendor fields and two PO line fields have it blank.
        if all(ws.cell(r, c).value in (None, "") for c in range(2, 13)):
            r += 1
            continue
        # "Vendor Address ID\n(On-demand ... Address Import Sync ID)" - the
        # parenthetical is a note about Concur's other importer, not the name.
        name = name.split("\n")[0].strip()
        fields.append((name,
                       str(req or "").split("\n")[0].strip(),
                       (str(proc).strip() if proc else "")))
        r += 1
    spec[key] = (rt, label, tab, fields)

EXPECTED = {"vendor_100": 4, "vendor_200": 62, "po_200": 61, "po_210": 20,
            "po_220": 20, "po_300": 52, "po_400": 42, "receipt_200": 29}
for k, n in EXPECTED.items():
    got = len(spec[k][3])
    assert got == n, f"{k}: spec workbook gave {got} fields, expected {n}"

lines = [
    '"""',
    "ERP_Concur_Spec - the field layout of every Concur record type in the three",
    "extract files, one entry per position.",
    "",
    "GENERATED - do not hand-edit. Emitted from",
    "~/Doppio/FMG/Concur_Record_Type_Specifications.xlsx (itself scraped from the",
    "SAP Help Portal, Concur Invoice Professional Edition Administration Guides,",
    f"version 2026_08) on {dt.date.today().isoformat()} by gen_spec.py. Re-run the",
    "generator if the workbook gains a record type or a field.",
    "",
    "Each record type is (record_type, label, source tab, [fields]) and each field",
    "is (name, required, proc column). `required` is the guide's own Y / N, kept",
    "verbatim - 'Y*' on PO 300 Expense Type means \"either this or Account Code,",
    "never both\", and 'N*' on a vendor field means \"required in some countries\".",
    "The section banners the workbook uses to group the vendor record ('Vendor",
    "Header', 'Vendor Address Detail') are dropped here: they are headings, not",
    "positions in the file.",
    '"""',
    "from __future__ import annotations",
    "",
    "",
    "# ---------------------------------------------------------------- layouts",
    "",
    "SPEC: dict[str, tuple[str, str, str, list[tuple[str, str, str]]]] = {",
]
for key, (rt, label, tab, fields) in spec.items():
    lines.append(f'    "{key}": ({rt!r}, {label!r}, {tab!r}, [')
    for name, req, proc in fields:
        lines.append(f"        ({name!r}, {req!r}, {proc!r}),")
    lines.append("    ]),")
lines += [
    "}",
    "",
    "# How wide each record is. A record that arrives with a different field count",
    "# is reported rather than guessed at - the whole file is positional, so one",
    "# extra comma shifts every field after it.",
    "WIDTHS: dict[str, int] = {k: len(v[3]) for k, v in SPEC.items()}",
    "",
    "",
    "def fields(key: str) -> list[tuple[str, str, str]]:",
    '    """The field list for a spec key, or an empty list if it is unknown."""',
    "    entry = SPEC.get(key)",
    "    return list(entry[3]) if entry else []",
    "",
    "",
    "def names(key: str) -> list[str]:",
    "    return [f[0] for f in fields(key)]",
    "",
    "",
    "def label(key: str) -> str:",
    "    entry = SPEC.get(key)",
    "    return entry[1] if entry else key",
    "",
    "",
    "def describe(key: str, values: list[str]) -> list[dict]:",
    '    """',
    "    Zip a record's values onto the spec, for the detail panel.",
    "",
    "    Positions the record does not reach come back empty, and positions the",
    "    spec does not know about (a file wider than the guide) come back named",
    '    "Field <n> (not in the spec)" rather than being dropped - an unexpected',
    "    field is exactly the thing worth seeing.",
    '    """',
    "    out = []",
    "    spec_fields = fields(key)",
    "    for i in range(max(len(spec_fields), len(values))):",
    "        if i < len(spec_fields):",
    "            name, req, proc = spec_fields[i]",
    "        else:",
    '            name, req, proc = f"Field {i + 1} (not in the spec)", "", ""',
    '        value = values[i] if i < len(values) else ""',
    "        out.append({",
    '            "position": i + 1, "name": name, "required": req,',
    '            "proc": proc, "value": value,',
    '            "padded": bool(value) and value != value.strip(),',
    "        })",
    "    return out",
    "",
]
OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print("wrote", OUT, OUT.stat().st_size, "bytes")
for k, v in spec.items():
    print(f"  {k:12s} rt={v[0]} width={len(v[3])}")
