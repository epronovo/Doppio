# MIG_MigrationReview.py
#
# Processes Excel workbooks that contain a [DTA_FIN_MVX] sheet: queries live
# M3 record counts via EXPORTMI, and produces a [migration review] copy with
# five new columns:
#   J  MT actual               - live count returned by the API
#   K  MT count vs actual      - =J{r}-F{r}
#   L  Precent Diff vs Actual  - =IFERROR(ABS(F-J)/((F+J)/2),0)  [0% format]
#   M  Table Description       - looked up from m3tables in the shared SQLite DB
#   N  Maintained By           - looked up from m3tables in the shared SQLite DB
#
# The original workbook is moved to "processed"; the reviewed copy lands in
# "output" with a _REVIEWED suffix.
#
# Ported to a library used by MIG_App.py: input()/tqdm are gone, tenant auth
# is an explicit MIG_Api.Tenant, and progress is reported through a job
# progress(done, total, message) callback instead of tqdm.write(). The
# repo-root config.py / InforMI.py dependency is gone too - see SQLITE_DB_PATH
# and the folder layout below.

from __future__ import annotations

import os
import shutil
import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import MIG_Api

# =============================================================================
# Shared SQLite DB (m3tables / m3TableCols) - also used by M3_Security and
# ADP_Concur. Populated by an unrelated tool; override with MIG_SYNC_DB.
# =============================================================================
SQLITE_DB_PATH = Path(os.environ.get("MIG_SYNC_DB", str(Path.home() / "sqlite" / "doppio.db")))

# =============================================================================
# Folder layout - package-local, NOT the shared ionapi/evs100 folders.
# =============================================================================
BASE_DIR = Path(__file__).parent.resolve()
INPUT_DIR = BASE_DIR / "input" / "mig_sync" / "migration_review"
PROCESSED_DIR = BASE_DIR / "output" / "mig_sync" / "migration_review" / "processed"
OUTPUT_DIR = BASE_DIR / "output" / "mig_sync" / "migration_review" / "reviewed"

REQUIRED_SHEET = "DTA_FIN_MVX"

# Column positions (1-based) for the five new columns
COL_J = 10   # MT actual
COL_K = 11   # MT count vs actual
COL_L = 12   # Precent Diff vs Actual
COL_M = 13   # Table Description
COL_N = 14   # Maintained By


def _noop_progress(done, total, message=""):
    pass


# =============================================================================
# Workbook helpers
# =============================================================================

def _find_header_row(ws) -> int | None:
    """Row number (1-based) whose col A = 'Library' and col B = 'Filename'."""
    for row in ws.iter_rows():
        if (str(row[0].value or "").strip().lower() == "library" and
                str(row[1].value or "").strip().lower() == "filename"):
            return row[0].row
    return None


def _collect_data_rows(ws, header_row: int) -> list[tuple[int, str]]:
    """(excel_row_number, filename) for every data row with a filename (col B)."""
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        filename = str(row[1].value or "").strip()
        if filename:
            rows.append((row[0].row, filename))
    return rows


# =============================================================================
# SQLite lookup - m3tables (Table Description / Maintained By)
# =============================================================================

def _lookup_table_info(filenames: list[str]) -> dict[str, tuple[str, str]]:
    """{filename: (tableDescription, tableMaintainedBy)} via one bulk query."""
    if not filenames:
        return {}

    placeholders = ",".join("?" * len(filenames))
    sql = (f"SELECT tablename, tableDescription, tableMaintainedBy "
           f"FROM m3tables WHERE tablename IN ({placeholders})")

    result: dict[str, tuple[str, str]] = {}
    try:
        with sqlite3.connect(SQLITE_DB_PATH) as conn:
            conn.row_factory = sqlite3.Row
            for row in conn.execute(sql, filenames):
                result[row["tablename"]] = (row["tableDescription"] or "",
                                             row["tableMaintainedBy"] or "")
    except Exception:
        pass  # DB unavailable - descriptions simply stay blank

    return result


# =============================================================================
# M3 API - EXPORTMI Select  (count(#) from <filename>)
# =============================================================================

def _get_live_count(tenant: MIG_Api.Tenant, filename: str,
                     session: requests.Session, max_retries: int = 3) -> int | str:
    """
    Calls EXPORTMI/Select and returns the integer count for filename, or
    empty string on unrecoverable error.
    """
    payload = {
        "program": "EXPORTMI",
        "transactions": [{"transaction": "Select",
                           "record": {"QERY": f"count(#) from {filename}"},
                           "selectedColumns": ["QERY", "REPL"]}],
    }
    for attempt in range(1, max_retries + 1):
        try:
            data = MIG_Api.post_to_m3(tenant, payload, session)
            records = data.get("results", [{}])[0].get("records", [])
            if records:
                return int(records[0].get("REPL", 0))
            return 0
        except Exception:
            if attempt == max_retries:
                return ""
            time.sleep(2)
    return ""


def _fetch_counts_parallel(tenant: MIG_Api.Tenant, data_rows: list[tuple[int, str]],
                            max_workers: int = 8,
                            progress=_noop_progress) -> dict[int, int | str]:
    """Fetch live counts for all rows in parallel. Returns {row_num: count}."""
    results: dict[int, int | str] = {}
    total = len(data_rows)
    done = 0
    progress(done, total, "Querying M3")

    with requests.Session() as session:
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_row = {
                executor.submit(_get_live_count, tenant, filename, session): row_num
                for row_num, filename in data_rows
            }
            for future in as_completed(future_to_row):
                row_num = future_to_row[future]
                results[row_num] = future.result()
                done += 1
                progress(done, total, "Querying M3")

    return results


# =============================================================================
# Sheet formatting helpers
# =============================================================================

_SUMMARY_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_SUMMARY_FONT = Font(bold=True, color="000000")
_OK_DIFF_FILL = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
_OK_DIFF_FONT = Font(bold=True, color="FFFFFF")
_QUESTIONABLE_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
_QUESTIONABLE_FONT = Font(bold=True, color="000000")


def _write_summary(ws, header_row: int, last_data_row: int) -> None:
    """
    Four-row summary block at K5:L8, driven by Remarks (col I) and MT count
    vs actual (col K):
      matching table counts - blank remarks AND K = 0
      protected             - non-blank remarks (any K)
      ok to be different     - blank remarks AND K > 0
      questionable           - blank remarks AND K < 0

    Rows matching "ok to be different" get their K cell filled green
    (_OK_DIFF_FILL); rows matching "questionable" get their K cell filled
    gold (_QUESTIONABLE_FILL) - applied per-row in process_workbook().
    """
    data_start = header_row + 1
    i_rng = f"I{data_start}:I{last_data_row}"
    k_rng = f"K{data_start}:K{last_data_row}"

    summary = [
        (5, "matching table counts:", f'=COUNTIFS({i_rng},"",{k_rng},0)'),
        (6, "protected:", f'=COUNTIF({i_rng},"<>")'),
        (7, "ok to be different:", f'=COUNTIFS({i_rng},"",{k_rng},">0")'),
        (8, "questionable:", f'=COUNTIFS({i_rng},"",{k_rng},"<0")'),
    ]

    for row_num, label, formula in summary:
        lbl = ws.cell(row=row_num, column=11)
        val = ws.cell(row=row_num, column=12)
        if row_num == 7:
            font, fill = _OK_DIFF_FONT, _OK_DIFF_FILL
        elif row_num == 8:
            font, fill = _QUESTIONABLE_FONT, _QUESTIONABLE_FILL
        else:
            font, fill = _SUMMARY_FONT, _SUMMARY_FILL
        lbl.value, lbl.font, lbl.fill = label, font, fill
        val.value, val.font, val.fill = formula, font, fill


def _reset_autofilter(ws, header_row: int, last_data_row: int, last_col: int) -> None:
    """Clear and re-apply the auto-filter across A{header_row}:<last_col><last_data_row>."""
    ws.auto_filter.ref = None
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_data_row}"


def _autofit_columns(ws, col_indices: list[int], header_row: int) -> None:
    """Width to fit each column's widest cell value, capped at 60 characters."""
    for col_idx in col_indices:
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row < header_row:
                continue
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


# =============================================================================
# Core workbook processor
# =============================================================================

def process_workbook(file_path: Path, tenant: MIG_Api.Tenant,
                      progress=_noop_progress) -> dict:
    """
    Process a single workbook. Returns a result dict:
        {"ok": bool, "message": str, "file": name, "output_path": str | None,
         "tables": int, "sqlite_matched": int}
    """
    file_path = Path(file_path)

    try:
        wb = load_workbook(file_path)
    except Exception as exc:
        return {"ok": False, "file": file_path.name, "message": f"Could not open workbook: {exc}"}

    if REQUIRED_SHEET not in wb.sheetnames:
        return {"ok": False, "file": file_path.name,
                "message": f"Sheet '{REQUIRED_SHEET}' not found - skipped."}

    ws = wb[REQUIRED_SHEET]
    header_row = _find_header_row(ws)
    if header_row is None:
        return {"ok": False, "file": file_path.name,
                "message": f"Header row (Library / Filename) not found in '{REQUIRED_SHEET}' - skipped."}

    ws.cell(row=header_row, column=COL_J).value = "MT actual"
    ws.cell(row=header_row, column=COL_K).value = "MT count vs actual"
    ws.cell(row=header_row, column=COL_L).value = "Precent Diff vs Actual"
    ws.cell(row=header_row, column=COL_M).value = "Table Description"
    ws.cell(row=header_row, column=COL_N).value = "Maintained By"

    data_rows = _collect_data_rows(ws, header_row)
    if not data_rows:
        return {"ok": False, "file": file_path.name, "message": "No data rows found - skipped."}

    counts = _fetch_counts_parallel(tenant, data_rows, progress=progress)
    unique_filenames = list({fn for _, fn in data_rows})
    table_info = _lookup_table_info(unique_filenames)

    for row_num, filename in data_rows:
        count = counts.get(row_num, "")
        description, maintained_by = table_info.get(filename, ("", ""))

        j_cell = ws.cell(row=row_num, column=COL_J)
        k_cell = ws.cell(row=row_num, column=COL_K)
        l_cell = ws.cell(row=row_num, column=COL_L)
        m_cell = ws.cell(row=row_num, column=COL_M)
        n_cell = ws.cell(row=row_num, column=COL_N)

        j_cell.value = count
        k_cell.value = f"=J{row_num}-F{row_num}"
        l_cell.value = f"=IFERROR(ABS(F{row_num} - J{row_num}) / ((F{row_num} + J{row_num}) / 2),0)"
        l_cell.number_format = "0%"
        m_cell.value = description
        n_cell.value = maintained_by

        remarks = ws.cell(row=row_num, column=9).value
        if not str(remarks or "").strip():
            try:
                diff = float(count) - float(ws.cell(row=row_num, column=6).value)
            except (TypeError, ValueError):
                diff = None
            if diff is not None and diff > 0:
                k_cell.fill = _OK_DIFF_FILL
                k_cell.font = _OK_DIFF_FONT
            elif diff is not None and diff < 0:
                k_cell.fill = _QUESTIONABLE_FILL
                k_cell.font = _QUESTIONABLE_FONT

    last_data_row = data_rows[-1][0] if data_rows else header_row
    _write_summary(ws, header_row, last_data_row)
    _reset_autofilter(ws, header_row, last_data_row, COL_N)
    _autofit_columns(ws, [COL_J, COL_K, COL_L, COL_M, COL_N], header_row)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

    output_name = file_path.stem + "_REVIEWED" + file_path.suffix
    output_path = OUTPUT_DIR / output_name
    try:
        wb.save(output_path)
    except Exception as exc:
        return {"ok": False, "file": file_path.name, "message": f"Could not save reviewed workbook: {exc}"}

    move_message = ""
    try:
        shutil.move(str(file_path), PROCESSED_DIR / file_path.name)
    except Exception as exc:
        move_message = f" (could not move original to processed: {exc})"

    return {
        "ok": True, "file": file_path.name, "output_path": str(output_path),
        "output_file": output_path.name, "tables": len(data_rows),
        "sqlite_matched": len(table_info), "sqlite_total": len(unique_filenames),
        "message": f"Saved -> {output_path.name}{move_message}",
    }


def list_pending_workbooks() -> list[Path]:
    """Every .xlsx sitting in INPUT_DIR (excluding Excel lock files)."""
    if not INPUT_DIR.is_dir():
        return []
    return sorted((f for f in INPUT_DIR.glob("*.xlsx") if not f.name.startswith("~$")),
                  key=lambda f: f.name.lower())
