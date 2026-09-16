# MIG_SyncPartnerMedia.py
"""
MIG_SyncPartnerMedia.py
-----------------------
Syncs partner media data (CRS949) from a SOURCE tenant to a DESTINATION
tenant via CRS949MI.LstPartnerMedia, CRS949MI.DltPartnerMedia, and
CRS949MI.AddPartnerEmail / AddPartnerMBM.

Steps
-----
1.  SOURCE tenant -> CRS949MI.LstPartnerMedia (SOURCE DIVI) -> records in memory.
2.  DEST   tenant -> CRS949MI.LstPartnerMedia (DEST   DIVI) -> records in memory.
3.  Diff: find records where any value field differs (matched on key fields),
    and records missing from DEST entirely.
4.  On DEST:
      a. CRS949MI.DltPartnerMedia + CRS949MI.AddPartnerEmail/AddPartnerMBM -
         update (delete & re-add) records that exist but differ.
      b. CRS949MI.AddPartnerEmail/AddPartnerMBM - add records missing from DEST.
5.  Write a summary .xlsx of all API calls made (always, even with nothing to do).

Notes
-----
Batch size is fixed at 100, matching the CLI (no prompt).

Ported to a library used by MIG_App.py: input()/print() are gone, tqdm is
replaced with a job progress(done, total, message) callback, and tenant auth
is an explicit MIG_Api.Tenant.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import MIG_Api

# =============================================================================
# Constants
# =============================================================================

PARTNER_MEDIA_COLS = [
    "DIVI", "DONR", "DOVA", "MEPF", "PRF1", "PRF2", "MEDC", "SEQN",
    "SIID", "MVIF", "METY", "FMTP", "COPY", "OUTP", "UDTA", "1UDT",
    "TFNO", "TFT1", "1TFT", "FLRN", "PAFD", "FSUX", "GNNM", "DEV",
    "FOVR", "BINX", "TOMA", "FRMA", "CCMA", "FAXT", "SUBJ", "NOTE",
    "CPPL", "LSID", "LSAD", "TEME", "E065", "SRD1", "SRD2", "RRD1",
    "RRD2", "RRD3", "TRAY", "LAYC", "MARI", "ARCH", "CNID", "EVPR",
    "BONM", "BOVB", "FIET", "EMBT", "EMGR", "FILM", "FNAM", "GRCO", "PRTP",
]

KEY_FIELDS = ["DIVI", "DONR", "DOVA", "MEPF", "PRF1", "PRF2", "MEDC", "SEQN"]
VALUE_FIELDS = [f for f in PARTNER_MEDIA_COLS if f not in KEY_FIELDS]

BATCH_SIZE = 100

# Maps MEDC value -> the CRS949MI Add transaction to use. Records whose MEDC
# is not in this map are skipped entirely (not deleted, not added).
MEDC_TRANSACTION: dict[str, str] = {
    "MAIL": "AddPartnerEmail",
    "MBMEVENT": "AddPartnerMBM",
}


def resolve_add_transaction(rec: dict) -> str | None:
    """The CRS949MI Add transaction for this record's MEDC, or None to skip."""
    return MEDC_TRANSACTION.get((rec.get("MEDC") or "").strip())


# =============================================================================
# CRS949MI.LstPartnerMedia -> list of dicts
# =============================================================================

def fetch_partner_media(tenant: MIG_Api.Tenant, session: requests.Session,
                         divi: str, label: str) -> list[dict]:
    """CRS949MI.LstPartnerMedia for the given division."""
    payload = {
        "program": "CRS949MI",
        "transactions": [{
            "transaction": "LstPartnerMedia",
            "record": {"DIVI": divi},
            "selectedColumns": PARTNER_MEDIA_COLS,
        }],
    }
    data = MIG_Api.post_to_m3(tenant, payload, session)

    rows: list[dict] = []
    for result in data.get("results", []):
        for record in result.get("records", []):
            rows.append({col: (record.get(col, "") or "").strip() for col in PARTNER_MEDIA_COLS})
    return rows


# =============================================================================
# Diff logic
# =============================================================================

def _row_key(row: dict) -> tuple:
    return tuple(row.get(f, "") for f in KEY_FIELDS)


def diff_partner_media(source_rows: list[dict],
                        dest_rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """
    Compare SOURCE and DEST partner-media lists.

    Returns (to_update, to_add):
        to_update - rows whose key exists in DEST but whose value fields differ
        to_add    - rows whose key is entirely absent from DEST
    """
    dest_index: dict[tuple, dict] = {_row_key(r): r for r in dest_rows}

    to_update: list[dict] = []
    to_add: list[dict] = []
    for src in source_rows:
        key = _row_key(src)
        dest = dest_index.get(key)
        if dest is None:
            to_add.append(src)
        elif any(src.get(f, "") != dest.get(f, "") for f in VALUE_FIELDS):
            to_update.append(src)

    return to_update, to_add


def preview_partner_media(to_update: list[dict], to_add: list[dict],
                           dest_rows: list[dict]) -> list[dict]:
    """
    Annotate the diff plan for display, exactly like the CLI's console table:
    ACTION (UPDATE/ADD/SKIP), TXN (the resolved CRS949MI transaction, or "-"
    for a skipped MEDC), CHANGES (a "FIELD: old->new" summary for updates).
    """
    dest_index: dict[tuple, dict] = {_row_key(r): r for r in dest_rows}

    def _changes_summary(src_row: dict) -> str:
        dest_row = dest_index.get(_row_key(src_row), {})
        parts = []
        for f in VALUE_FIELDS:
            sv, dv = src_row.get(f, ""), dest_row.get(f, "")
            if sv != dv:
                parts.append(f'{f}: "{dv}"->"{sv}"')
        return " | ".join(parts)

    def _annotate(row: dict, base_action: str, changes: str) -> dict:
        txn = resolve_add_transaction(row)
        return {**row, "ACTION": "SKIP" if txn is None else base_action,
                "TXN": "-" if txn is None else txn, "CHANGES": changes}

    return ([_annotate(row, "UPDATE", _changes_summary(row)) for row in to_update]
            + [_annotate(row, "ADD", "") for row in to_add])


# =============================================================================
# Build API records
# =============================================================================

def build_dlt_record(dest_row: dict, dest_divi: str) -> dict:
    """Record for CRS949MI.DltPartnerMedia using DEST values + DEST DIVI."""
    rec = {col: dest_row.get(col, "") for col in PARTNER_MEDIA_COLS}
    rec["DIVI"] = dest_divi
    return {k: v for k, v in rec.items() if v != ""}


def build_add_record(src_row: dict, dest_divi: str) -> dict:
    """Record for CRS949MI.Add* using SOURCE values + DEST DIVI."""
    rec = {col: src_row.get(col, "") for col in PARTNER_MEDIA_COLS}
    rec["DIVI"] = dest_divi
    return {k: v for k, v in rec.items() if v != ""}


# =============================================================================
# CRS949MI.DltPartnerMedia on DEST (batched)
# =============================================================================

def _noop_progress(done, total, message=""):
    pass


def run_dlt_batched(tenant: MIG_Api.Tenant, records: list[dict],
                     session: requests.Session, batch_size: int,
                     progress=_noop_progress) -> tuple[list[dict], list[tuple[dict, str]]]:
    """
    Deletes existing DEST records (before re-adding updated SOURCE values).
    Returns (successes, failures) where failures is [(record, error), ...].
    """
    successes: list[dict] = []
    failures: list[tuple[dict, str]] = []
    total = len(records)
    progress(0, total, "CRS949MI.DltPartnerMedia")

    for i in range(0, total, batch_size):
        batch = records[i:i + batch_size]
        payload = {
            "program": "CRS949MI",
            "transactions": [{"transaction": "DltPartnerMedia", "record": rec,
                               "selectedColumns": PARTNER_MEDIA_COLS} for rec in batch],
        }
        try:
            result = MIG_Api.post_to_m3(tenant, payload, session)
            api_results = result.get("results", [])
            for j, res in enumerate(api_results):
                err = (res.get("errorMessage") or "").strip() if isinstance(res, dict) else ""
                rec = batch[j] if j < len(batch) else {}
                if err:
                    failures.append((rec, err))
                else:
                    successes.append(rec)
            for k in range(len(api_results), len(batch)):
                failures.append((batch[k], "No result returned"))
        except Exception as exc:
            for rec in batch:
                failures.append((rec, str(exc)))
        progress(min(i + batch_size, total), total, "CRS949MI.DltPartnerMedia")

    return successes, failures


# =============================================================================
# CRS949MI.Add* on DEST (batched, MEDC-routed)
# =============================================================================

def run_add_batched(tenant: MIG_Api.Tenant, records: list[dict],
                     session: requests.Session, batch_size: int,
                     progress=_noop_progress, desc: str = "add"
                     ) -> tuple[list[dict], list[tuple[dict, str]], list[dict]]:
    """
    Adds records to DEST, routing each record to the correct transaction based
    on its MEDC value (MAIL -> AddPartnerEmail, MBMEVENT -> AddPartnerMBM,
    anything else -> skipped, not sent to the API).

    Returns (successes, failures, skipped).
    """
    groups: dict[str, list[dict]] = {}
    skipped: list[dict] = []
    for rec in records:
        txn = resolve_add_transaction(rec)
        if txn is None:
            skipped.append(rec)
        else:
            groups.setdefault(txn, []).append(rec)

    successes: list[dict] = []
    failures: list[tuple[dict, str]] = []
    total = sum(len(g) for g in groups.values())
    done = 0
    progress(0, total, f"CRS949MI Add [{desc}]")

    for txn, grp in groups.items():
        for i in range(0, len(grp), batch_size):
            batch = grp[i:i + batch_size]
            payload = {
                "program": "CRS949MI",
                "transactions": [{"transaction": txn, "record": rec,
                                   "selectedColumns": PARTNER_MEDIA_COLS} for rec in batch],
            }
            try:
                result = MIG_Api.post_to_m3(tenant, payload, session)
                api_results = result.get("results", [])
                for j, res in enumerate(api_results):
                    err = (res.get("errorMessage") or "").strip() if isinstance(res, dict) else ""
                    rec = batch[j] if j < len(batch) else {}
                    if err:
                        failures.append((rec, err))
                    else:
                        successes.append(rec)
                for k in range(len(api_results), len(batch)):
                    failures.append((batch[k], "No result returned"))
            except Exception as exc:
                for rec in batch:
                    failures.append((rec, str(exc)))
            done += len(batch)
            progress(done, total, f"CRS949MI.{txn} [{desc}]")

    return successes, failures, skipped


# =============================================================================
# XLSX summary (always written, not just on errors)
# =============================================================================

def export_summary_xlsx(
    source_count: int, dest_count: int,
    dlt_successes: list[dict], dlt_failures: list[tuple[dict, str]],
    upd_add_successes: list[dict], upd_add_failures: list[tuple[dict, str]],
    upd_skipped: list[dict],
    add_successes: list[dict], add_failures: list[tuple[dict, str]],
    add_skipped: list[dict],
    out_dir: Path,
) -> Path:
    """
    Always writes an xlsx summarising all API calls:
      Sheet 1 - Summary          : row per step with success/failed/skipped counts
      Sheet 2 - DltPartnerMedia  : detail rows for every Dlt result
      Sheet 3 - Add (Update)     : detail rows for re-adds after Dlt
      Sheet 4 - Add (New)        : detail rows for net-new adds
      Sheet 5 - Skipped          : records excluded due to unsupported MEDC
    """
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"MIG_SyncPartnerMedia_Summary_{ts}.xlsx"

    wb = Workbook()

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", start_color="2F5496")
    ok_fill = PatternFill("solid", start_color="E2EFDA")
    err_fill = PatternFill("solid", start_color="FCE4D6")
    skip_fill = PatternFill("solid", start_color="FFF2CC")
    info_fill = PatternFill("solid", start_color="DDEBF7")
    ctr = Alignment(horizontal="center")
    arial = Font(name="Arial")

    def _hdr(ws, col, row, val):
        c = ws.cell(row=row, column=col, value=val)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = ctr
        return c

    def _cell(ws, col, row, val, fill=None, align=None):
        c = ws.cell(row=row, column=col, value=val)
        c.font = arial
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        return c

    def _write_detail_sheet(ws, ok_list, failures, cols, extra_col=None, extra_fn=None):
        full_cols = cols + ([extra_col] if extra_col else []) + ["STATUS"]
        for ci, f in enumerate(full_cols, start=1):
            _hdr(ws, ci, 1, f)
        all_rows = [(rec, "OK") for rec in ok_list] + [(rec, err) for rec, err in failures]
        for ri, (rec, status) in enumerate(all_rows, start=2):
            fill = ok_fill if status == "OK" else err_fill
            for ci, f in enumerate(full_cols, start=1):
                if f == "STATUS":
                    val = status
                elif f == extra_col and extra_fn:
                    val = extra_fn(rec)
                else:
                    val = rec.get(f, "")
                _cell(ws, ci, ri, val, fill=fill)
        for ci, f in enumerate(full_cols, start=1):
            sample = [(status if f == "STATUS" else extra_fn(rec) if (f == extra_col and extra_fn)
                       else rec.get(f, "")) for rec, status in all_rows]
            max_w = max(len(str(v)) for v in [f] + sample) if all_rows else len(f)
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(max_w + 2, 60)

    def _write_skipped_sheet(ws, skipped):
        full_cols = PARTNER_MEDIA_COLS + ["REASON"]
        for ci, f in enumerate(full_cols, start=1):
            _hdr(ws, ci, 1, f)
        for ri, rec in enumerate(skipped, start=2):
            for ci, f in enumerate(full_cols, start=1):
                val = f"MEDC '{rec.get('MEDC','')}' not in MEDC_TRANSACTION" if f == "REASON" else rec.get(f, "")
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = arial
                c.fill = skip_fill
        for ci, f in enumerate(full_cols, start=1):
            sample = [(f"MEDC '{r.get('MEDC','')}' not in MEDC_TRANSACTION" if f == "REASON" else r.get(f, ""))
                      for r in skipped]
            max_w = max(len(str(v)) for v in [f] + sample) if skipped else len(f)
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(max_w + 2, 60)

    ws_sum = wb.active
    ws_sum.title = "Summary"
    for ci, heading in enumerate(["Step", "Transaction(s)", "Success", "Failed", "Skipped", "Total"], start=1):
        _hdr(ws_sum, ci, 1, heading)

    steps = [
        ("1 - Extract SOURCE", "LstPartnerMedia", source_count, 0, 0),
        ("2 - Extract DEST", "LstPartnerMedia", dest_count, 0, 0),
        ("3a - Delete (update)", "DltPartnerMedia", len(dlt_successes), len(dlt_failures), 0),
        ("3b - Re-add (update)", "AddPartnerEmail / AddPartnerMBM", len(upd_add_successes), len(upd_add_failures), len(upd_skipped)),
        ("4  - Add (new)", "AddPartnerEmail / AddPartnerMBM", len(add_successes), len(add_failures), len(add_skipped)),
    ]
    for ri, (step_label, txn, ok, fail, skip) in enumerate(steps, start=2):
        fill = err_fill if fail > 0 else (skip_fill if skip > 0 else info_fill)
        _cell(ws_sum, 1, ri, step_label, fill=fill)
        _cell(ws_sum, 2, ri, txn, fill=fill)
        _cell(ws_sum, 3, ri, ok, fill=fill, align=ctr)
        _cell(ws_sum, 4, ri, fail, fill=fill, align=ctr)
        _cell(ws_sum, 5, ri, skip, fill=fill, align=ctr)
        _cell(ws_sum, 6, ri, ok + fail + skip, fill=fill, align=ctr)

    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 38
    for col in "CDEF":
        ws_sum.column_dimensions[col].width = 12

    if dlt_successes or dlt_failures:
        _write_detail_sheet(wb.create_sheet("DltPartnerMedia"), dlt_successes, dlt_failures, PARTNER_MEDIA_COLS)

    if upd_add_successes or upd_add_failures:
        _write_detail_sheet(wb.create_sheet("Add (Update)"), upd_add_successes, upd_add_failures,
                             PARTNER_MEDIA_COLS, extra_col="TRANSACTION",
                             extra_fn=lambda r: resolve_add_transaction(r) or "-")

    if add_successes or add_failures:
        _write_detail_sheet(wb.create_sheet("Add (New)"), add_successes, add_failures,
                             PARTNER_MEDIA_COLS, extra_col="TRANSACTION",
                             extra_fn=lambda r: resolve_add_transaction(r) or "-")

    all_skipped_recs = upd_skipped + add_skipped
    if all_skipped_recs:
        _write_skipped_sheet(wb.create_sheet("Skipped"), all_skipped_recs)

    wb.save(out_path)
    return out_path
