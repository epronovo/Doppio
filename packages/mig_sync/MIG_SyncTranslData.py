# MIG_SyncTranslData.py
"""
MIG_SyncTranslData.py
----------------------
Extracts translation data (MBMTRN / MBMTRD) from a SOURCE tenant and writes
it to an EVS100-compatible Excel file for import into a DESTINATION tenant,
plus (optionally) pushes it directly into DEST via CRS881MI.

Steps
-----
1.  SOURCE tenant -> EXPORTMI.Select (MBMTRN) -> header rows.
2.                -> EXPORTMI.Select (MBMTRD) -> detail rows.
3.  Join header + details on IDTR.
4.  Export to an EVS100-format Excel file
    (evs100/ToProcess/API_CRS881MI_<timestamp>.xlsx) with:
      - Control sheet listing both data sheets
      - Sheet "API_CRS881MI_AddTranslation" - one row per unique header
      - Sheet "API_CRS881MI_AddTranslData"  - one row per detail
5.  Optionally call CRS881MI.AddTranslation / AddTranslData directly on DEST.
6.  Optionally upload the file to M3 via the File Management REST API (PUT)
    and trigger processing via EVS100MI.ImportFile.

Translation data is global - not scoped to a company or division - so both
tenants here use MIG_Api.Tenant(global_scope=True), which drops &cono=/&divi=
from the API URL entirely (see MIG_Api.Tenant.api_url). This replaces the
original CLI's private _apply_global_url()/_get_ion_token_global()/_global_url()
trick with the same effect through a cleaner mechanism.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import requests
import xlsxwriter
from openpyxl import Workbook as OWorkbook
from openpyxl.styles import Alignment, Font, PatternFill

import MIG_Api

# =============================================================================
# Constants
# =============================================================================

EXPORTMI_SEP = "^"
DEFAULT_BATCH_SIZE = 100

# MBMTRN - Translation header
EXPORTMI_HDR_QUERY = (
    "TRIDTR,TRTRQF,TRMSTD,TRMVRS,TRBMSG,TRIBOB,TRELMP,TRELMD,TRELMC,TRMBMC from MBMTRN"
)
# MBMTRD - Translation details
EXPORTMI_DTL_QUERY = "TDDIVI,TDIDTR,TDTX15,TDMVXP,TDEXTP,TDMVXD,TDMBMD,TDTX40 from MBMTRD"

# Mapping from MBMTRN column names -> CRS881MI API field names.
HDR_FIELD_MAP: dict[str, str] = {
    "TRIDTR": "IDTR", "TRTRQF": "TRQF", "TRMSTD": "MSTD", "TRMVRS": "MVRS",
    "TRBMSG": "BMSG", "TRIBOB": "IBOB", "TRELMP": "ELMP", "TRELMD": "ELMD",
    "TRELMC": "ELMC", "TRMBMC": "MBMC",
}

# Mapping from MBMTRD column names -> CRS881MI API field names.
# TDIDTR is the join key only; it is NOT passed to the API.
DTL_FIELD_MAP: dict[str, str] = {
    "TDDIVI": "DIVI", "TDTX15": "TX15", "TDMVXP": "MVXP", "TDEXTP": "EXTP",
    "TDMVXD": "MVXD", "TDMBMD": "MBMD", "TDTX40": "TX40",
}

ADD_TRANSLATION_FIELDS = list(HDR_FIELD_MAP.values())
ADD_TRANSL_DATA_FIELDS = ["CONO"] + ADD_TRANSLATION_FIELDS + list(DTL_FIELD_MAP.values())


def _noop_progress(done, total, message=""):
    pass


# =============================================================================
# EXPORTMI.Select -> parse REPL rows
# =============================================================================

def fetch_exportmi(tenant: MIG_Api.Tenant, session: requests.Session,
                    query: str, label: str) -> list[dict]:
    """
    Calls EXPORTMI.Select with the given query and returns a list of dicts,
    one per data row, keyed by the column names found in the HDRS row.
    """
    payload = {
        "program": "EXPORTMI",
        "transactions": [{
            "transaction": "Select",
            "record": {"QERY": query, "SEPC": EXPORTMI_SEP, "HDRS": "1"},
            "selectedColumns": ["QERY", "SEPC", "HDRS", "REPL"],
        }],
    }
    data = MIG_Api.post_to_m3(tenant, payload, session)

    repl_rows: list[str] = []
    for result in data.get("results", []):
        for record in result.get("records", []):
            repl_val = record.get("REPL", "")
            if repl_val:
                repl_rows.append(repl_val)

    if not repl_rows:
        return []

    col_names = [c.strip() for c in repl_rows[0].rstrip(EXPORTMI_SEP).split(EXPORTMI_SEP)]

    parsed: list[dict] = []
    for raw in repl_rows[1:]:
        values = raw.rstrip(EXPORTMI_SEP).split(EXPORTMI_SEP)
        values += [""] * (len(col_names) - len(values))
        row = {col_names[i]: values[i].strip() for i in range(len(col_names))}
        parsed.append(row)

    return parsed


# =============================================================================
# Join MBMTRN header + MBMTRD details on IDTR
# =============================================================================

def join_header_details(hdr_rows: list[dict], dtl_rows: list[dict]) -> list[dict]:
    """
    Joins MBMTRN header rows and MBMTRD detail rows on IDTR (one-to-many).
    Orphaned detail rows (no matching header) are silently skipped.
    """
    hdr_index: dict[str, dict] = {r["TRIDTR"]: r for r in hdr_rows}

    combined: list[dict] = []
    for dtl in dtl_rows:
        hdr = hdr_index.get(dtl.get("TDIDTR", ""))
        if hdr is not None:
            combined.append({**hdr, **dtl})
    return combined


# =============================================================================
# Build API records from combined rows
# =============================================================================

def build_trn_record(hdr_row: dict) -> dict:
    """CRS881MI.AddTranslation record from a header row."""
    rec: dict = {}
    for src_key, api_key in HDR_FIELD_MAP.items():
        val = hdr_row.get(src_key, "")
        if val:
            rec[api_key] = val
    return rec


def build_trd_record(combined_row: dict, company: str) -> dict:
    """CRS881MI.AddTranslData record from a combined header + detail row."""
    rec: dict = {"CONO": company}
    for src_key, api_key in {**HDR_FIELD_MAP, **DTL_FIELD_MAP}.items():
        val = combined_row.get(src_key, "")
        if val:
            rec[api_key] = val
    return rec


def build_records(hdr_rows: list[dict], dtl_rows: list[dict],
                   company: str) -> tuple[list[dict], list[dict]]:
    """
    Step 3: join header + details, dedupe headers by TRIDTR, and build the
    two API record sets. Returns (trn_records, trd_records).
    """
    combined = join_header_details(hdr_rows, dtl_rows)

    seen_idtr: set[str] = set()
    trn_records: list[dict] = []
    for row in hdr_rows:
        idtr = row.get("TRIDTR", "")
        if idtr not in seen_idtr:
            seen_idtr.add(idtr)
            trn_records.append(build_trn_record(row))

    trd_records = [build_trd_record(row, company) for row in combined]
    return trn_records, trd_records


# =============================================================================
# EVS100 Excel export
# =============================================================================

def export_evs100_xlsx(trn_records: list[dict], trd_records: list[dict],
                        out_dir: Path) -> Path:
    """
    Writes records to an EVS100 import file with two data sheets:
    API_CRS881MI_AddTranslation (headers) and API_CRS881MI_AddTranslData
    (details). Returns the path of the written file.
    """
    SHEET_TRN = "API_CRS881MI_AddTranslation"
    SHEET_TRD = "API_CRS881MI_AddTranslData"

    DESCS_TRN = [None, "Translation ID", "Qualifier", "Message standard", "Message version",
                 "Business message", "In/Out", "Element path",
                 "Element description", "Element container", "Message context"]
    REQD_TRN = ["no"] + ["yes"] * len(ADD_TRANSLATION_FIELDS)
    COLS_TRN = ["MESSAGE"] + ADD_TRANSLATION_FIELDS

    DESCS_TRD = [None, "Company", "Translation ID", "Qualifier", "Message standard",
                 "Message version", "Business message", "In/Out", "Element path",
                 "Element description", "Element container", "Message context",
                 "Division", "Text 15", "MVX path", "Extension type", "MVX description",
                 "Message description", "Text 40"]
    REQD_TRD = ["no"] + ["yes"] * len(ADD_TRANSL_DATA_FIELDS)
    COLS_TRD = ["MESSAGE"] + ADD_TRANSL_DATA_FIELDS

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"API_CRS881MI_{ts}.xlsx"

    wb = xlsxwriter.Workbook(str(out_path))

    ws_ctrl = wb.add_worksheet("Control")
    for col, val in enumerate(["Worksheet", "Description", "Data"]):
        ws_ctrl.write(0, col, val)
    for col, val in enumerate([SHEET_TRN, "Translation header", "x"]):
        ws_ctrl.write(1, col, val)
    for col, val in enumerate([SHEET_TRD, "Translation details", "x"]):
        ws_ctrl.write(2, col, val)

    def _write_data_sheet(ws, cols, descs, reqd, records, fields):
        for col, name in enumerate(cols):
            ws.write(0, col, name)
        for col, desc in enumerate(descs):
            if desc is not None:
                ws.write(1, col, desc)
        for col, req in enumerate(reqd):
            ws.write(2, col, req)
        for row_idx, rec in enumerate(records, start=3):
            for col, field_name in enumerate(fields, start=1):
                val = rec.get(field_name, "")
                if val:
                    ws.write(row_idx, col, val)

    ws_trn = wb.add_worksheet(SHEET_TRN)
    _write_data_sheet(ws_trn, COLS_TRN, DESCS_TRN, REQD_TRN, trn_records, ADD_TRANSLATION_FIELDS)

    ws_trd = wb.add_worksheet(SHEET_TRD)
    _write_data_sheet(ws_trd, COLS_TRD, DESCS_TRD, REQD_TRD, trd_records, ADD_TRANSL_DATA_FIELDS)

    wb.close()
    return out_path


# =============================================================================
# Direct API calls to DEST (mirrors MIG_SyncPartnerRef pattern)
# =============================================================================

def _run_batched(tenant: MIG_Api.Tenant, transaction: str, fields: list[str],
                  records: list[dict], session: requests.Session, batch_size: int,
                  progress) -> tuple[list[dict], list[tuple[dict, str]]]:
    successes: list[dict] = []
    failures: list[tuple[dict, str]] = []
    total = len(records)
    label = f"CRS881MI.{transaction}"
    progress(0, total, label)

    for i in range(0, total, batch_size):
        batch = records[i:i + batch_size]
        payload = {
            "program": "CRS881MI",
            "transactions": [{"transaction": transaction, "record": rec,
                               "selectedColumns": fields} for rec in batch],
        }
        try:
            result = MIG_Api.post_to_m3(tenant, payload, session)
            api_results = result.get("results", [])
            for j, res in enumerate(api_results):
                err = (res.get("errorMessage") or "").strip() if isinstance(res, dict) else ""
                rec = batch[j] if j < len(batch) else {}
                if err:
                    failures.append((rec, err))
                else:
                    successes.append(rec)
            for k in range(len(api_results), len(batch)):
                failures.append((batch[k], "No result returned"))
        except Exception as exc:
            for rec in batch:
                failures.append((rec, str(exc)))
        progress(min(i + batch_size, total), total, label)

    return successes, failures


def run_add_translation_batched(tenant: MIG_Api.Tenant, records: list[dict],
                                 session: requests.Session, batch_size: int,
                                 progress=_noop_progress) -> tuple[list[dict], list[tuple[dict, str]]]:
    """CRS881MI.AddTranslation on DEST for each record (batched)."""
    return _run_batched(tenant, "AddTranslation", ADD_TRANSLATION_FIELDS, records,
                         session, batch_size, progress)


def run_add_transl_data_batched(tenant: MIG_Api.Tenant, records: list[dict],
                                 session: requests.Session, batch_size: int,
                                 progress=_noop_progress) -> tuple[list[dict], list[tuple[dict, str]]]:
    """CRS881MI.AddTranslData on DEST for each record (batched)."""
    return _run_batched(tenant, "AddTranslData", ADD_TRANSL_DATA_FIELDS, records,
                         session, batch_size, progress)


def export_api_errors_xlsx(
    trn_successes: list[dict], trn_failures: list[tuple[dict, str]],
    trd_successes: list[dict], trd_failures: list[tuple[dict, str]],
    out_dir: Path,
) -> Path:
    """
    Sheets: Summary (error-message totals for both steps), AddTranslation Err,
    AddTranslData Err (one row per failed record). Returns the written path.
    """
    from collections import Counter

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"MIG_SyncTranslData_Errors_{ts}.xlsx"

    wb = OWorkbook()
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", start_color="2F5496")
    ok_fill = PatternFill("solid", start_color="E2EFDA")
    err_fill = PatternFill("solid", start_color="FCE4D6")
    ctr = Alignment(horizontal="center")

    def _write_error_sheet(ws, fields, failures):
        for col, f in enumerate(fields, start=1):
            c = ws.cell(row=1, column=col, value=f)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = ctr
        for row_i, (rec, err) in enumerate(failures, start=2):
            for col, f in enumerate(fields, start=1):
                val = err if f == "ERROR" else rec.get(f, "")
                c = ws.cell(row=row_i, column=col, value=val)
                c.font = Font(name="Arial")
        for col, f in enumerate(fields, start=1):
            max_len = max(len(f), *(len(str(rec.get(f, "") if f != "ERROR" else err))
                                     for rec, err in failures))
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max_len + 2, 60)

    ws_sum = wb.active
    ws_sum.title = "Summary"
    for col, heading in enumerate(["Step", "Message", "Count"], start=1):
        c = ws_sum.cell(row=1, column=col, value=heading)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = ctr

    row_i = 2
    for step_label, successes, failures in [
        ("AddTranslation", trn_successes, trn_failures),
        ("AddTranslData", trd_successes, trd_failures),
    ]:
        counts: Counter = Counter()
        counts["OK"] = len(successes)
        for _, err in failures:
            counts[err] += 1
        for msg, count in counts.most_common():
            fill = ok_fill if msg == "OK" else err_fill
            for col, val in enumerate([step_label, msg, count], start=1):
                c = ws_sum.cell(row=row_i, column=col, value=val)
                c.fill = fill
                c.font = Font(name="Arial")
                if col == 3:
                    c.alignment = ctr
            row_i += 1

    ws_sum.column_dimensions["A"].width = 20
    ws_sum.column_dimensions["B"].width = 80
    ws_sum.column_dimensions["C"].width = 10

    if trn_failures:
        _write_error_sheet(wb.create_sheet("AddTranslation Err"), ADD_TRANSLATION_FIELDS + ["ERROR"], trn_failures)
    if trd_failures:
        _write_error_sheet(wb.create_sheet("AddTranslData Err"), ADD_TRANSL_DATA_FIELDS + ["ERROR"], trd_failures)

    wb.save(out_path)
    return out_path
